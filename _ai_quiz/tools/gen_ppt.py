# -*- coding: utf-8 -*-
"""基于模板生成《站前小学一年级新生家长会——政教处分享》PPT（6页）。

模板: 2025522学校家长会（改）.pptx（标题页 + 结尾页）
本脚本:
  1. 修改模板标题页文本 → 新封面
  2. 复制模板设计元素（背景/白卡/三张角图）新建 4 个内容页
  3. 修改模板结尾页文本 → 总结页
"""
import os
import copy
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

TPL = r"f:\桌面\桌面所有图标\工作复盘\教师\工作\2025522学校家长会（改）.pptx"
OUT_DIR = r"F:\cursor20260624\output"
OUT = os.path.join(OUT_DIR, "站前小学一年级新生家长会_政教处分享.pptx")

# ---------- 配色 / 字体 ----------
KAITI = "楷体"
GREEN_DARK = RGBColor(0x21, 0x33, 0x14)   # 深绿（正文/标题主色，同模板文字色）
GREEN_TEAL = RGBColor(0x38, 0x5F, 0x52)   # 模板背景绿
GOLD = RGBColor(0xC9, 0xA0, 0x63)         # 金色点缀
LIGHT_GREEN = RGBColor(0xE9, 0xF1, 0xEC)  # 浅绿填充
LIGHT_GOLD = RGBColor(0xFB, 0xF3, 0xE4)   # 浅金填充
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def set_run_font(r, size, color, bold, typeface=KAITI):
    f = r.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = color
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", typeface)


def clear_tf(tf):
    tf.clear()


def fill_tf_lines(tf, lines, size, color, bold=True, typeface=KAITI,
                  align=PP_ALIGN.CENTER, line_spacing=1.0, space_after=0):
    """lines: 每段为 str / (text,size,color,bold) / 多个 run 组成的列表"""
    clear_tf(tf)
    first = True
    for ln in lines:
        if isinstance(ln, str):
            runs_spec = [(ln, size, color, bold)]
        elif isinstance(ln, tuple):
            runs_spec = [ln]
        else:
            runs_spec = ln
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        if line_spacing:
            p.line_spacing = line_spacing
        if space_after:
            p.space_after = Pt(space_after)
        for (text, sz, col, bd) in runs_spec:
            r = p.add_run()
            r.text = text
            set_run_font(r, sz, col, bd, typeface)


def style_shape_text(sp, x, y, w, h, lines, size, color, bold=True,
                     align=PP_ALIGN.CENTER, line_spacing=1.1, space_after=4,
                     margins=(0.12, 0.12, 0.05, 0.05), anchor=MSO_ANCHOR.MIDDLE):
    tf = sp.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(margins[0])
    tf.margin_right = Inches(margins[1])
    tf.margin_top = Inches(margins[2])
    tf.margin_bottom = Inches(margins[3])
    tf.vertical_anchor = anchor
    fill_tf_lines(tf, lines, size, color, bold, align=align,
                  line_spacing=line_spacing, space_after=space_after)
    return sp


def add_pill(slide, x, y, w, h, text, fill=LIGHT_GREEN, border=GOLD, size=17):
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid()
    sp.fill.fore_color.rgb = fill
    sp.line.color.rgb = border
    sp.line.width = Pt(1.2)
    sp.shadow.inherit = False
    style_shape_text(sp, x, y, w, h, [text], size, GREEN_DARK, True)
    return sp


def add_box(slide, x, y, w, h, fill, border, lines, size=18,
            align=PP_ALIGN.CENTER, line_spacing=1.15, space_after=4, bold=True):
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid()
    sp.fill.fore_color.rgb = fill
    sp.line.color.rgb = border
    sp.line.width = Pt(1.3)
    sp.shadow.inherit = False
    style_shape_text(sp, x, y, w, h, lines, size, GREEN_DARK, bold,
                     align=align, line_spacing=line_spacing, space_after=space_after)
    return sp


def add_textbox(slide, x, y, w, h, lines, size, color=GREEN_DARK, bold=True,
                align=PP_ALIGN.LEFT, line_spacing=1.15, space_after=8):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    fill_tf_lines(tf, lines, size, color, bold, align=align,
                  line_spacing=line_spacing, space_after=space_after)
    return tb


def add_bullets(slide, x, y, w, items, size=19, gap=14, marker="◆",
                mcolor=GOLD, tcolor=GREEN_DARK, line_spacing=1.1):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(2))
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for item in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = PP_ALIGN.LEFT
        p.line_spacing = line_spacing
        p.space_after = Pt(gap)
        r = p.add_run()
        r.text = marker + "  "
        set_run_font(r, size, mcolor, True)
        r2 = p.add_run()
        r2.text = item
        set_run_font(r2, size, tcolor, False)
    return tb


def copy_design_shapes(src_slide, dst_slide, names):
    """把模板 slide 中的设计元素(矩形/图片)深拷贝到新 slide。"""
    for sh in src_slide.shapes:
        if sh.name in names:
            el = copy.deepcopy(sh._element)
            if sh.shape_type == 13:  # PICTURE
                blip = el.find(".//{%s}blip" % A)
                rid = blip.get("{%s}embed" % R_NS)
                part = src_slide.part.related_part(rid)
                nrid = dst_slide.part.relate_to(part, RT.IMAGE)
                blip.set("{%s}embed" % R_NS, nrid)
            dst_slide.shapes._spTree.append(el)
            names.discard(sh.name)


def remove_placeholders(slide):
    for sh in list(slide.shapes):
        if sh.is_placeholder:
            sh._element.getparent().remove(sh._element)


# ============================================================
prs = Presentation(TPL)
s_title = prs.slides[0]
s_end = prs.slides[1]

# ---------- 1) 标题页：替换文本 ----------
title_shape = None
sub_shape = None
foot_shape = None
for sh in s_title.shapes:
    if sh.name == "文本框 10":
        title_shape = sh
    elif sh.name == "文本框 5":
        sub_shape = sh
    elif sh.name == "文本框 4":
        foot_shape = sh

# 主标题：两行（适当上移，避免与副标题重叠）
title_shape.top = Inches(1.9)
fill_tf_lines(title_shape.text_frame,
              [("扣好人生第一粒扣子", 46, RGBColor(0xB2, 0x5E, 0x25), True),
               ("携手共育新成长", 46, RGBColor(0xB2, 0x5E, 0x25), True)],
              size=46, color=RGBColor(0xB2, 0x5E, 0x25), bold=True,
              align=PP_ALIGN.CENTER, line_spacing=1.0, space_after=2)
# 副标题
fill_tf_lines(sub_shape.text_frame,
              ["站前小学一年级新生家长会——政教处分享"],
              size=34, color=GREEN_DARK, bold=True,
              align=PP_ALIGN.CENTER, line_spacing=1.0, space_after=0)
# 落款
fill_tf_lines(foot_shape.text_frame,
              ["金华市站前小学德育中心"],
              size=24, color=GREEN_DARK, bold=True,
              align=PP_ALIGN.CENTER, line_spacing=1.0, space_after=0)

# ---------- 2) 四个内容页 ----------
blank = prs.slide_masters[0].slide_layouts[6]  # 空白
DESIGN = {"矩形 11", "矩形 1", "图片 17", "图片 21", "图片 3"}


def new_content_slide():
    s = prs.slides.add_slide(blank)
    remove_placeholders(s)
    copy_design_shapes(s_title, s, set(DESIGN))
    return s


def add_title(slide, text):
    # 标题
    add_textbox(slide, 1.1, 0.72, 11.13, 0.62, [text], 30, GREEN_DARK, True,
                align=PP_ALIGN.CENTER, space_after=0)
    # 金色下划线
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                 Inches(5.87), Inches(1.42), Inches(1.6), Inches(0.045))
    bar.fill.solid()
    bar.fill.fore_color.rgb = GOLD
    bar.line.fill.background()
    bar.shadow.inherit = False


# --- 第2页：站前十礼 ---
s2 = new_content_slide()
add_title(s2, "践行站前十礼，养成良好校园品行")
RITES = ["入学礼", "问候礼", "同学礼", "师生礼", "排队礼",
         "上课礼", "课间礼", "集会礼", "就餐礼", "放学礼"]
pw, gap, ph = 2.0, 0.22, 0.72
xs = 1.22
for i in range(5):
    add_pill(s2, xs + i * (pw + gap), 1.95, pw, ph, RITES[i])
for i in range(5):
    add_pill(s2, xs + i * (pw + gap), 2.82, pw, ph, RITES[5 + i])
add_bullets(s2, 1.4, 4.05, 10.6,
            ["覆盖学生在校一日校园生活，把礼仪转化为日常行为习惯",
             "家校同频同向，共同助力孩子文明礼仪养成"],
            size=20, gap=16)

# --- 第3页：同伴小矛盾 ---
s3 = new_content_slide()
add_title(s3, "理性看待孩子之间的同伴小矛盾")
add_bullets(s3, 1.4, 1.78, 10.6, [
    "低段孩子相处，磕磕碰碰、发生小矛盾属于成长常态",
    "孩子复述事件容易带入主观感受，容易断章取义、只讲部分事实",
    "无意的触碰，孩子可能表述为“打”；无心举动容易被当成故意行为",
    "只听孩子单方面讲述，容易产生误会",
], size=19, gap=12)
add_box(s3, 1.4, 4.05, 10.6, 0.7, LIGHT_GREEN, GOLD,
        [[("✅ 建议：", 19, GOLD, True),
          ("先接纳情绪，不急定性对错，及时与班主任沟通核实完整情况", 19, GREEN_DARK, False)]],
        size=19, align=PP_ALIGN.LEFT, line_spacing=1.1, space_after=0, bold=False)
add_box(s3, 1.4, 4.9, 10.6, 0.7, LIGHT_GOLD, GOLD,
        [("💡 同伴相处，也是落实“同学礼”，是孩子重要的成长必修课", 19, GREEN_DARK, False)],
        size=19, align=PP_ALIGN.LEFT, line_spacing=1.1, space_after=0, bold=False)

# --- 第4页：家庭规则 ---
s4 = new_content_slide()
add_title(s4, "家庭养育：建立简单清晰的规则")
add_box(s4, 1.6, 1.78, 10.1, 1.05, LIGHT_GOLD, GOLD,
        [("核心金句", 18, GOLD, True),
         ("规则只有一条，在规定的时间做好规定的事。", 22, GREEN_DARK, True)],
        size=22, line_spacing=1.15, space_after=3)
add_bullets(s4, 1.4, 3.35, 10.6, [
    "建立家庭稳定作息节奏：起床、学习、玩耍、就寝时间明确",
    "家庭规则贵在少而坚定，不用堆砌繁杂禁令",
    "稳定的时间边界，带给孩子充足内心安全感",
], size=20, gap=18)

# --- 第5页：亲子陪伴三件小事 ---
s5 = new_content_slide()
add_title(s5, "坚持三件小事，给予孩子情感滋养")
cards = [
    ("✅ 1个拥抱", "传递接纳与爱意，给予孩子心理底气，缓解入学焦虑"),
    ("✅ 10分钟睡前按摩", "轻柔肢体安抚，消解在校疲惫，舒缓情绪压力"),
    ("✅ 10分钟闲聊", "放下说教，平等聊天，倾听孩子校园见闻与烦恼"),
]
cw, cgap, ch, cy = 3.4, 0.32, 2.6, 2.2
cx = 1.15
for i, (t, d) in enumerate(cards):
    x = cx + i * (cw + cgap)
    sp = add_box(s5, x, cy, cw, ch, LIGHT_GREEN, GOLD,
                 [(t, 21, GREEN_DARK, True),
                  (d, 16.5, GREEN_TEAL, False)],
                 size=21, align=PP_ALIGN.CENTER, line_spacing=1.25,
                 space_after=8, bold=True)
add_textbox(s5, 1.4, 5.15, 10.6, 0.6,
            ["闲聊，也是了解孩子同伴交往真实状态的重要窗口"],
            18, GREEN_TEAL, False, align=PP_ALIGN.CENTER, space_after=0)

# ---------- 3) 结尾页：修改 slide2 ----------
for sh in s_end.shapes:
    if sh.name == "文本框 5":      # 顶部小字：与标题重叠，删掉
        sh._element.getparent().remove(sh._element)
    elif sh.name == "文本框 6":    # 大标题 → 移到左侧列
        sh.text_frame.word_wrap = True
        sh.left = Inches(1.3)
        sh.top = Inches(1.7)
        sh.width = Inches(6.9)
        sh.height = Inches(1.5)
        fill_tf_lines(sh.text_frame,
                      [("家校同心，陪伴孩子", 36, GREEN_DARK, True),
                       ("平稳开启小学时光", 36, GREEN_DARK, True)],
                      size=36, color=GREEN_DARK, bold=True,
                      align=PP_ALIGN.CENTER, line_spacing=1.1, space_after=2)
    elif sh.name == "文本框 4":    # 底部日期 → 落款
        fill_tf_lines(sh.text_frame, ["金华市站前小学德育中心"],
                      size=22, color=GREEN_DARK, bold=True,
                      align=PP_ALIGN.CENTER, line_spacing=1.0, space_after=0)
    elif sh.name == "文本框 7":    # 主讲人 → 与要点重叠，删掉
        sh._element.getparent().remove(sh._element)

# 结尾页三要点（左侧列）
add_bullets(s_end, 1.3, 3.5, 6.9, [
    "一年级重点：规则意识、文明礼仪、情绪安全感、同伴交往能力",
    "学校抓习惯养成，家庭守好家庭规则，严慈相济",
    "理性看待同伴磕碰，家校互信互通，携手陪伴孩子走好每一步",
], size=18, gap=12)

# ---------- 4) 调整顺序：标题(1) 内容(2-5) 结尾(6) ----------
sldIdLst = prs.slides._sldIdLst
ids = list(sldIdLst)          # [s_title, s_end, c1, c2, c3, c4]
sldIdLst.remove(ids[1])       # 拿出结尾页
sldIdLst.append(ids[1])       # 放到最后

# ---------- 保存 ----------
os.makedirs(OUT_DIR, exist_ok=True)
prs.save(OUT)
print("saved:", OUT)
