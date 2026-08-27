# -*- coding: utf-8 -*-
"""Inspect a PPTX template: slides, layouts, shape texts."""
import sys
from pptx import Presentation
from pptx.util import Emu

path = r"f:\桌面\桌面所有图标\工作复盘\教师\工作\2025522学校家长会（改）.pptx"
prs = Presentation(path)

print("Slide size:", prs.slide_width, "x", prs.slide_height, "=", Emu(prs.slide_width).inches, "x", Emu(prs.slide_height).inches, "inches")
print("Num slides:", len(prs.slides))
print("Num layouts:", len(prs.slide_masters[0].slide_layouts) if prs.slide_masters else 0)
print("=" * 80)

for i, layout in enumerate(prs.slide_masters[0].slide_layouts):
    print(f"[Layout {i}] name={layout.name!r}")

print("=" * 80)
for idx, slide in enumerate(prs.slides):
    print(f"\n=== Slide {idx+1} === layout={slide.slide_layout.name!r}")
    for sh in slide.shapes:
        t = ""
        if sh.has_text_frame:
            t = " | ".join(p.text for p in sh.text_frame.paragraphs if p.text.strip())
        print(f"  - {sh.shape_type}, name={sh.name!r}, pos=({Emu(sh.left).inches:.2f},{Emu(sh.top).inches:.2f}) "
              f"size=({Emu(sh.width).inches:.2f}x{Emu(sh.height).inches:.2f})")
        if t:
            print(f"      text: {t[:200]}")
